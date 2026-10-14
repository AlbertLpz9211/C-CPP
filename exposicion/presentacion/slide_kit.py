"""Utilidades de maquetacion para la presentacion de la exposicion.

Envuelve python-pptx con un conjunto reducido de funciones (portada, seccion,
vinetas, bloque de codigo, tabla, cita, demostracion) para que el archivo de
contenido se lea como un guion y no como codigo de dibujo.

El resultado se guarda como .pptx y se convierte a .odp con LibreOffice.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

# --- Paleta ----------------------------------------------------------------
FONDO       = RGBColor(0xFA, 0xFA, 0xF7)
FONDO_OSC   = RGBColor(0x14, 0x26, 0x38)
TITULO      = RGBColor(0x14, 0x26, 0x38)
TEXTO       = RGBColor(0x2E, 0x3A, 0x46)
TENUE       = RGBColor(0x74, 0x82, 0x8C)
ACENTO      = RGBColor(0x0B, 0x6E, 0x6E)
ALERTA      = RGBColor(0xA8, 0x43, 0x1B)
BLANCO      = RGBColor(0xFF, 0xFF, 0xFF)
CODIGO_FONDO = RGBColor(0xF0, 0xF2, 0xF1)
CODIGO_TEXTO = RGBColor(0x1E, 0x2B, 0x33)
CODIGO_COM   = RGBColor(0x6B, 0x7A, 0x74)

FUENTE = "Helvetica Neue"
MONO = "Menlo"

ANCHO = 13.333
ALTO = 7.5
MARGEN = 0.85


class Mazo:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Inches(ANCHO)
        self.prs.slide_height = Inches(ALTO)
        self._n = 0

    # -- infraestructura ----------------------------------------------------
    def _nueva(self, oscura=False):
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        fondo = s.background.fill
        fondo.solid()
        fondo.fore_color.rgb = FONDO_OSC if oscura else FONDO
        self._n += 1
        return s

    def _caja(self, s, x, y, w, h):
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        return tf

    def _rico(self, parrafo, texto, size, color, fuente=FUENTE, bold=False):
        """Escribe texto admitiendo **negritas** y `codigo` en linea."""
        import re
        piezas = re.split(r"(\*\*.+?\*\*|`.+?`)", texto)
        for pieza in piezas:
            if not pieza:
                continue
            r = parrafo.add_run()
            if pieza.startswith("**") and pieza.endswith("**"):
                r.text = pieza[2:-2]
                r.font.bold = True
                r.font.name = fuente
                r.font.color.rgb = TITULO if color == TEXTO else color
            elif pieza.startswith("`") and pieza.endswith("`"):
                r.text = pieza[1:-1]
                r.font.name = MONO
                r.font.color.rgb = ACENTO
                r.font.size = Pt(size - 1)
                continue
            else:
                r.text = pieza
                r.font.name = fuente
                r.font.bold = bold
                r.font.color.rgb = color
            r.font.size = Pt(size)

    def _notas(self, s, texto):
        if texto:
            s.notes_slide.notes_text_frame.text = texto.strip()

    def _numero(self, s):
        tf = self._caja(s, ANCHO - 1.35, ALTO - 0.62, 0.9, 0.3)
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.RIGHT
        r = p.add_run()
        r.text = str(self._n)
        r.font.size = Pt(11)
        r.font.name = FUENTE
        r.font.color.rgb = TENUE

    def _titulo(self, s, texto, sub=None):
        tf = self._caja(s, MARGEN, 0.55, ANCHO - 2 * MARGEN, 1.0)
        p = tf.paragraphs[0]
        self._rico(p, texto, 30, TITULO, bold=True)
        for r in p.runs:
            r.font.bold = True
        linea = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGEN), Inches(1.32),
                                   Inches(1.5), Inches(0.045))
        linea.fill.solid()
        linea.fill.fore_color.rgb = ACENTO
        linea.line.fill.background()
        linea.shadow.inherit = False
        if sub:
            tfs = self._caja(s, MARGEN, 1.48, ANCHO - 2 * MARGEN, 0.45)
            ps = tfs.paragraphs[0]
            self._rico(ps, sub, 15, TENUE)

    # -- tipos de diapositiva ----------------------------------------------
    def portada(self, titulo, subtitulo, autor, pie, notas=None):
        s = self._nueva(oscura=True)
        barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(MARGEN), Inches(2.35),
                                   Inches(2.0), Inches(0.06))
        barra.fill.solid()
        barra.fill.fore_color.rgb = ACENTO
        barra.line.fill.background()
        barra.shadow.inherit = False

        tf = self._caja(s, MARGEN, 2.7, ANCHO - 2 * MARGEN, 1.8)
        p = tf.paragraphs[0]
        self._rico(p, titulo, 44, BLANCO, bold=True)

        tf2 = self._caja(s, MARGEN, 4.35, ANCHO - 2 * MARGEN, 0.6)
        self._rico(tf2.paragraphs[0], subtitulo, 20, RGBColor(0x9F, 0xC4, 0xC4))

        tf3 = self._caja(s, MARGEN, 5.6, ANCHO - 2 * MARGEN, 1.0)
        p3 = tf3.paragraphs[0]
        self._rico(p3, autor, 17, BLANCO)
        p4 = tf3.add_paragraph()
        p4.space_before = Pt(6)
        self._rico(p4, pie, 14, RGBColor(0x8E, 0x9E, 0xAB))
        self._notas(s, notas)
        return s

    def seccion(self, numero, titulo, minutos, puntos=None, notas=None):
        s = self._nueva(oscura=True)
        tf = self._caja(s, MARGEN, 2.1, 2.0, 1.4)
        r = tf.paragraphs[0].add_run()
        r.text = numero
        r.font.size = Pt(72)
        r.font.bold = True
        r.font.name = FUENTE
        r.font.color.rgb = ACENTO

        tf2 = self._caja(s, MARGEN, 3.35, ANCHO - 2 * MARGEN, 1.2)
        self._rico(tf2.paragraphs[0], titulo, 38, BLANCO, bold=True)
        for r in tf2.paragraphs[0].runs:
            r.font.bold = True

        tf3 = self._caja(s, MARGEN, 4.6, ANCHO - 2 * MARGEN, 0.5)
        self._rico(tf3.paragraphs[0], minutos, 16, RGBColor(0x9F, 0xC4, 0xC4))

        if puntos:
            tf4 = self._caja(s, MARGEN, 5.25, ANCHO - 2 * MARGEN, 1.5)
            for i, punto in enumerate(puntos):
                p = tf4.paragraphs[0] if i == 0 else tf4.add_paragraph()
                p.space_before = Pt(5)
                self._rico(p, "·  " + punto, 15, RGBColor(0xC5, 0xD0, 0xD8))
        self._notas(s, notas)
        return s

    def vinetas(self, titulo, items, sub=None, notas=None, size=19, x=None, w=None, y=2.0):
        s = self._nueva()
        self._titulo(s, titulo, sub)
        x = MARGEN if x is None else x
        w = (ANCHO - 2 * MARGEN) if w is None else w
        tf = self._caja(s, x, y, w, ALTO - y - 0.8)
        primero = True
        for item in items:
            nivel = 0
            texto = item
            if isinstance(item, tuple):
                nivel, texto = item
            p = tf.paragraphs[0] if primero else tf.add_paragraph()
            primero = False
            p.space_before = Pt(4 if nivel else 11)
            if texto == "":
                continue
            vineta = "—  " if nivel else "•  "
            tam = size - 3 if nivel else size
            r = p.add_run()
            r.text = vineta
            r.font.size = Pt(tam)
            r.font.name = FUENTE
            r.font.color.rgb = ACENTO if not nivel else TENUE
            self._rico(p, texto, tam, TEXTO)
            pPr = p._p.get_or_add_pPr()
            if nivel:
                p.level = 1
                pPr.set("marL", str(int(Inches(0.78))))
            else:
                pPr.set("marL", str(int(Inches(0.33))))
            pPr.set("indent", str(-int(Inches(0.33))))
        self._notas(s, notas)
        self._numero(s)
        return s

    def _dibuja_codigo(self, s, codigo, x, y, w, h, size=14, resaltar=()):
        caja = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y),
                                  Inches(w), Inches(h))
        caja.fill.solid()
        caja.fill.fore_color.rgb = CODIGO_FONDO
        caja.line.color.rgb = RGBColor(0xDD, 0xE2, 0xE0)
        caja.line.width = Pt(0.75)
        caja.shadow.inherit = False
        caja.adjustments[0] = 0.03

        barra = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y + 0.08),
                                   Inches(0.055), Inches(h - 0.16))
        barra.fill.solid()
        barra.fill.fore_color.rgb = ACENTO
        barra.line.fill.background()
        barra.shadow.inherit = False

        tf = self._caja(s, x + 0.28, y + 0.16, w - 0.45, h - 0.3)
        lineas = codigo.strip("\n").split("\n")
        for i, linea in enumerate(lineas):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(0)
            p.line_spacing = 1.06
            destacada = i in resaltar
            desnuda = linea.strip()
            if desnuda.startswith("//"):
                trozos = [(linea, CODIGO_COM)]
            elif "//" in linea:
                cortada = linea.index("//")
                trozos = [(linea[:cortada], CODIGO_TEXTO), (linea[cortada:], CODIGO_COM)]
            else:
                trozos = [(linea, CODIGO_TEXTO)]
            for texto, color in trozos:
                r = p.add_run()
                r.text = texto if texto else " "
                r.font.name = MONO
                r.font.size = Pt(size)
                r.font.color.rgb = ALERTA if destacada else color
                r.font.bold = destacada

    @staticmethod
    def _alto_codigo(codigo, size, disponible):
        lineas = len(codigo.strip("\n").split("\n"))
        return min(disponible, lineas * size * 0.0175 + 0.42)

    def codigo(self, titulo, codigo, sub=None, pie=None, notas=None, size=14, resaltar=()):
        s = self._nueva()
        self._titulo(s, titulo, sub)
        disponible = ALTO - 2.05 - (0.75 if pie else 0.45)
        alto = self._alto_codigo(codigo, size, disponible)
        self._dibuja_codigo(s, codigo, MARGEN, 2.0, ANCHO - 2 * MARGEN, alto, size, resaltar)
        if pie:
            tf = self._caja(s, MARGEN, 2.0 + alto + 0.22, ANCHO - 2 * MARGEN, 0.6)
            self._rico(tf.paragraphs[0], pie, 16, ACENTO)
        self._notas(s, notas)
        self._numero(s)
        return s

    def dos_codigos(self, titulo, izq_tit, izq, der_tit, der, sub=None, pie=None,
                    notas=None, size=13):
        s = self._nueva()
        self._titulo(s, titulo, sub)
        w = (ANCHO - 2 * MARGEN - 0.4) / 2
        disponible = ALTO - 2.55 - (0.7 if pie else 0.3)
        alto = max(self._alto_codigo(izq, size, disponible),
                   self._alto_codigo(der, size, disponible))
        for i, (rot, cod) in enumerate([(izq_tit, izq), (der_tit, der)]):
            x = MARGEN + i * (w + 0.4)
            tf = self._caja(s, x, 1.98, w, 0.35)
            self._rico(tf.paragraphs[0], rot, 15, ALERTA if i == 0 else ACENTO, bold=True)
            for r in tf.paragraphs[0].runs:
                r.font.bold = True
            self._dibuja_codigo(s, cod, x, 2.4, w, alto, size)
        if pie:
            tf = self._caja(s, MARGEN, 2.4 + alto + 0.2, ANCHO - 2 * MARGEN, 0.6)
            self._rico(tf.paragraphs[0], pie, 16, ACENTO)
        self._notas(s, notas)
        self._numero(s)
        return s

    def tabla(self, titulo, cabeceras, filas, sub=None, notas=None, anchos=None,
              size=14, pie=None):
        s = self._nueva()
        self._titulo(s, titulo, sub)
        nf = len(filas) + 1
        nc = len(cabeceras)
        alto_disp = ALTO - 2.05 - (0.8 if pie else 0.45)
        alto = min(alto_disp, 0.42 * nf + 0.15)
        forma = s.shapes.add_table(nf, nc, Inches(MARGEN), Inches(2.0),
                                   Inches(ANCHO - 2 * MARGEN), Inches(alto))
        t = forma.table
        t.first_row = True
        if anchos:
            total = sum(anchos)
            disp = ANCHO - 2 * MARGEN
            for j, a in enumerate(anchos):
                t.columns[j].width = Inches(disp * a / total)
        for j, texto in enumerate(cabeceras):
            celda = t.cell(0, j)
            celda.text = ""
            celda.fill.solid()
            celda.fill.fore_color.rgb = FONDO_OSC
            celda.vertical_anchor = MSO_ANCHOR.MIDDLE
            celda.margin_left = celda.margin_right = Inches(0.12)
            p = celda.text_frame.paragraphs[0]
            mono = texto.startswith("`") and texto.endswith("`")
            r = p.add_run()
            r.text = texto[1:-1] if mono else texto.replace("**", "")
            r.font.bold = True
            r.font.size = Pt(size - 1 if mono else size)
            r.font.name = MONO if mono else FUENTE
            r.font.color.rgb = BLANCO
        for i, fila in enumerate(filas, start=1):
            for j, texto in enumerate(fila):
                celda = t.cell(i, j)
                celda.text = ""
                celda.fill.solid()
                celda.fill.fore_color.rgb = BLANCO if i % 2 else RGBColor(0xF2, 0xF4, 0xF3)
                celda.vertical_anchor = MSO_ANCHOR.MIDDLE
                celda.margin_left = celda.margin_right = Inches(0.12)
                p = celda.text_frame.paragraphs[0]
                mono = texto.startswith("`") and texto.endswith("`")
                if mono:
                    r = p.add_run()
                    r.text = texto[1:-1]
                    r.font.name = MONO
                    r.font.size = Pt(size - 1)
                    r.font.color.rgb = CODIGO_TEXTO
                else:
                    self._rico(p, texto, size, TEXTO)
        if pie:
            tf = self._caja(s, MARGEN, 2.0 + alto + 0.25, ANCHO - 2 * MARGEN, 0.6)
            self._rico(tf.paragraphs[0], pie, 15, ACENTO)
        self._notas(s, notas)
        self._numero(s)
        return s

    def mensaje(self, texto, fuente=None, notas=None, size=34):
        s = self._nueva(oscura=True)
        tf = self._caja(s, 1.3, 2.4, ANCHO - 2.6, 2.6)
        p = tf.paragraphs[0]
        p.line_spacing = 1.25
        self._rico(p, texto, size, BLANCO)
        if fuente:
            tf2 = self._caja(s, 1.3, 5.3, ANCHO - 2.6, 0.6)
            self._rico(tf2.paragraphs[0], fuente, 15, RGBColor(0x9F, 0xC4, 0xC4))
        self._notas(s, notas)
        return s

    def demo(self, titulo, comando, puntos, notas=None):
        s = self._nueva()
        self._titulo(s, "Demostración en vivo", titulo)
        self._dibuja_codigo(s, comando, MARGEN, 2.05, ANCHO - 2 * MARGEN,
                            0.55 + 0.28 * comando.strip().count("\n"), 15)
        y = 2.05 + 0.55 + 0.28 * comando.strip().count("\n") + 0.4
        tf = self._caja(s, MARGEN, y, ANCHO - 2 * MARGEN, ALTO - y - 0.7)
        for i, punto in enumerate(puntos):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.space_before = Pt(10)
            r = p.add_run()
            r.text = "→  "
            r.font.size = Pt(18)
            r.font.name = FUENTE
            r.font.color.rgb = ACENTO
            self._rico(p, punto, 18, TEXTO)
        self._notas(s, notas)
        self._numero(s)
        return s

    def practica(self, titulo, archivo, tareas, minutos, notas=None):
        s = self._nueva(oscura=True)
        tf = self._caja(s, MARGEN, 1.5, ANCHO - 2 * MARGEN, 0.8)
        self._rico(tf.paragraphs[0], titulo, 36, BLANCO, bold=True)
        for r in tf.paragraphs[0].runs:
            r.font.bold = True
        tf2 = self._caja(s, MARGEN, 2.45, ANCHO - 2 * MARGEN, 0.45)
        self._rico(tf2.paragraphs[0], minutos + "   ·   " + archivo, 15,
                   RGBColor(0x9F, 0xC4, 0xC4))
        tf3 = self._caja(s, MARGEN, 3.3, ANCHO - 2 * MARGEN, 3.4)
        for i, tarea in enumerate(tareas):
            p = tf3.paragraphs[0] if i == 0 else tf3.add_paragraph()
            p.space_before = Pt(12)
            r = p.add_run()
            r.text = "%d.  " % (i + 1)
            r.font.size = Pt(19)
            r.font.name = FUENTE
            r.font.bold = True
            r.font.color.rgb = ACENTO
            self._rico(p, tarea, 19, RGBColor(0xE4, 0xEA, 0xEE))
        self._notas(s, notas)
        return s

    def guardar(self, ruta):
        self.prs.save(ruta)
        return ruta
